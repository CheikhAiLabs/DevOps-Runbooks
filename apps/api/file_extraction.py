"""Bounded local file readers, isolated from the model-serving process.

Add a reader to EXTRACTORS to support another format. Readers never execute
uploaded code, formulas, macros, or external links, or write to supplied paths.
"""
from __future__ import annotations

import asyncio
import contextlib
import csv
import io
import json
import logging
import os
import signal
import sys
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILES = 5
MAX_EXTRACTED_CHARS = 12000
MAX_FILE_CONTEXT_CHARS = 20000
MAX_PDF_PAGES = 50
MAX_SHEETS = 5
MAX_ROWS = 200
MAX_COLUMNS = 30
MAX_IMAGE_PIXELS = 16_000_000
EXTRACTION_TIMEOUT = 90
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".log", ".json", ".yaml", ".yml", ".xml",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".sh", ".bash", ".hcl", ".tf",
    ".tfvars", ".ini", ".conf", ".toml", ".sql", ".go", ".java", ".c",
    ".cpp", ".h", ".rs",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


class ExtractionError(Exception):
    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status


@dataclass
class ExtractedText:
    parts: list[str] = field(default_factory=list)
    characters: int = 0
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)

    def append(self, text: str) -> bool:
        text = text.replace("\x00", "").strip()
        if not text:
            return True
        remaining = MAX_EXTRACTED_CHARS - self.characters
        piece = ("\n" if self.parts else "") + text
        self.parts.append(piece[:remaining])
        self.characters += min(len(piece), remaining)
        if len(piece) > remaining:
            self.truncated = True
        return self.characters < MAX_EXTRACTED_CHARS

    @property
    def text(self) -> str:
        return "".join(self.parts)


def safe_name(name: str) -> str:
    name = name.replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(c for c in name if c.isprintable()).strip()
    if not name or len(name) > 200:
        raise ExtractionError(400, "Use a filename between 1 and 200 characters.")
    return name


def decode_text(raw: bytes) -> str:
    from charset_normalizer import from_bytes

    # Try BOM encodings first, including UTF-32 before its UTF-16 prefix.
    for marker, encoding in (
        (b"\xff\xfe\x00\x00", "utf-32"), (b"\x00\x00\xfe\xff", "utf-32"),
        (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16"),
        (b"\xef\xbb\xbf", "utf-8-sig"),
    ):
        if raw.startswith(marker):
            try:
                text = raw.decode(encoding)
            except UnicodeError as exc:
                raise ExtractionError(422, "The file contains invalid encoded text.") from exc
            break
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeError:
            match = from_bytes(raw[:65536]).best()
            if match is None:
                raise ExtractionError(422, "The text encoding could not be detected.")
            try:
                text = raw.decode(match.encoding)
            except UnicodeError as exc:
                raise ExtractionError(422, "The file contains mixed or invalid text encoding.") from exc

    controls = sum(not c.isprintable() and c not in "\n\r\t\f" for c in text)
    if "\x00" in text or controls > max(2, len(text) // 100):
        raise ExtractionError(415, "This file appears to contain binary data, not text.")
    return text


def read_text(raw: bytes, result: ExtractedText) -> None:
    # Config, XML and JSON stay as text; no entities or uploaded code are evaluated.
    result.append(decode_text(raw))


def read_csv(raw: bytes, result: ExtractedText) -> None:
    text = decode_text(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    csv.field_size_limit(MAX_FILE_BYTES)
    for number, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        if number > MAX_ROWS:
            result.truncated = True
            break
        if len(row) > MAX_COLUMNS:
            result.truncated = True
        cells = " | ".join(json.dumps(cell, ensure_ascii=False) for cell in row[:MAX_COLUMNS])
        if not result.append(f"Row {number}: {cells}"):
            result.truncated = True
            break


def read_image(raw: bytes, result: ExtractedText) -> None:
    import pytesseract
    from PIL import Image, ImageOps

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(raw)) as original:
            if original.format not in {"PNG", "JPEG", "WEBP"}:
                raise ExtractionError(415, "Only PNG, JPEG and WEBP images are supported.")
            if original.width * original.height > MAX_IMAGE_PIXELS:
                raise ExtractionError(413, "Images must contain at most 16 million pixels.")
            with ImageOps.exif_transpose(original).convert("RGB") as image:
                try:
                    text = pytesseract.image_to_string(image, lang="eng+fra", timeout=60)
                except pytesseract.TesseractNotFoundError as exc:
                    raise ExtractionError(503, "Image reading is temporarily unavailable.") from exc
                except RuntimeError as exc:
                    raise ExtractionError(422, "Image reading failed or took too long.") from exc
    result.append(text)
    if not result.text.strip():
        result.warnings.append("No readable text was found in this image.")


def read_pdf(raw: bytes, result: ExtractedText) -> None:
    import pymupdf

    if not raw[:1024].lstrip().startswith(b"%PDF-"):
        raise ExtractionError(415, "The file is not a valid PDF.")
    with pymupdf.open(stream=raw, filetype="pdf") as document:
        if document.needs_pass:
            raise ExtractionError(422, "Password-protected PDFs are not supported.")
        result.truncated = document.page_count > MAX_PDF_PAGES
        empty_pages = 0
        for index in range(min(document.page_count, MAX_PDF_PAGES)):
            text = document[index].get_text("text", sort=True).strip()
            if not text:
                empty_pages += 1
                continue
            if not result.append(f"Page {index + 1}\n{text}"):
                result.truncated = True
                break
        if empty_pages:
            result.warnings.append("Some PDF pages have no native text; scanned pages are not read.")
        if not result.text.strip():
            result.warnings.append("No native text was found. Scanned PDFs are not supported yet.")


def check_office_archive(raw: bytes, required: str) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if required not in archive.namelist():
            raise ExtractionError(415, "The file does not match its Office document type.")
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 32 * 1024 * 1024:
            raise ExtractionError(413, "The expanded document exceeds the 32 MB processing limit.")
        if any(item.flag_bits & 1 for item in entries):
            raise ExtractionError(422, "Encrypted Office documents are not supported.")


def read_docx(raw: bytes, result: ExtractedText) -> None:
    from docx import Document
    from docx.text.paragraph import Paragraph

    check_office_archive(raw, "word/document.xml")
    document = Document(io.BytesIO(raw))
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if not result.append(block.text):
                result.truncated = True
                break
        else:
            for index, row in enumerate(block.rows):
                if index >= MAX_ROWS:
                    result.truncated = True
                    break
                if len(row.cells) > MAX_COLUMNS:
                    result.truncated = True
                if not result.append(" | ".join(cell.text for cell in row.cells[:MAX_COLUMNS])):
                    result.truncated = True
                    return
    result.warnings.append("Document body and tables only; embedded images and objects are not read.")


def read_xlsx(raw: bytes, result: ExtractedText) -> None:
    from openpyxl import load_workbook

    check_office_archive(raw, "xl/workbook.xml")
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        result.truncated = len(sheets) > MAX_SHEETS
        for sheet in sheets[:MAX_SHEETS]:
            heading_written = False
            # Ignore unreliable dimensions; the iteration itself stays bounded.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLUMNS + 1, values_only=True)
            for index, row in enumerate(rows, 1):
                if index > MAX_ROWS:
                    result.truncated = True
                    break
                if row[MAX_COLUMNS] is not None:
                    result.truncated = True
                cells = list(row[:MAX_COLUMNS])
                while cells and cells[-1] is None:
                    cells.pop()
                if not cells:
                    continue
                if not heading_written:
                    if not result.append(f"Sheet: {sheet.title}"):
                        result.truncated = True
                        return
                    heading_written = True
                line = " | ".join("" if value is None else str(value) for value in cells)
                if not result.append(f"Row {index}: {line}"):
                    result.truncated = True
                    return
    finally:
        workbook.close()
        result.warnings.append("Visible sheets only; formulas use saved values and are not recalculated.")


# The registry is the single backend extension point; MIME headers are not trusted.
EXTRACTORS: dict[str, tuple[str, Callable[[bytes, ExtractedText], None]]] = {
    **{extension: ("text", read_text) for extension in TEXT_EXTENSIONS},
    **{extension: ("image", read_image) for extension in IMAGE_EXTENSIONS},
    ".csv": ("spreadsheet", read_csv),
    ".pdf": ("pdf", read_pdf),
    ".docx": ("document", read_docx),
    ".xlsx": ("spreadsheet", read_xlsx),
}


def extract_file(raw: bytes, filename: str) -> dict:
    name = safe_name(filename)
    extension = Path(name).suffix.lower()
    if extension not in EXTRACTORS:
        raise ExtractionError(415, "This file type is not supported.")
    if not raw:
        raise ExtractionError(400, "The file is empty.")
    if len(raw) > MAX_FILE_BYTES:
        raise ExtractionError(413, "Files must be 10 MB or smaller.")

    kind, reader = EXTRACTORS[extension]
    result = ExtractedText()
    reader(raw, result)
    if result.truncated:
        result.warnings.append("Only part of this file was extracted because of processing limits.")
    if not result.text.strip() and not result.warnings:
        result.warnings.append("No readable content was found in this file.")
    return {
        "name": name, "size": len(raw), "kind": kind, "text": result.text,
        "characters": len(result.text), "truncated": result.truncated,
        "warnings": result.warnings,
    }


async def extract_isolated(raw: bytes, filename: str) -> dict:
    """Kill timed-out/cancelled parsers and their OCR children without loading models."""
    process = await asyncio.create_subprocess_exec(
        sys.executable, str(Path(__file__).resolve()), filename,
        stdin=asyncio.subprocess.PIPE, stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.DEVNULL, start_new_session=True,
        env={**os.environ, "OMP_NUM_THREADS": "1", "OPENBLAS_NUM_THREADS": "1", "MKL_NUM_THREADS": "1"},
    )
    try:
        stdout, _ = await asyncio.wait_for(process.communicate(raw), EXTRACTION_TIMEOUT)
        if process.returncode:
            raise ExtractionError(422, "The file exceeded processing limits or could not be read.")
        payload = json.loads(stdout)
        if "error" in payload:
            raise ExtractionError(payload["status"], payload["error"])
        return payload
    except asyncio.TimeoutError as exc:
        raise ExtractionError(504, "File processing timed out.") from exc
    finally:
        # Even when the parser exited, an OCR child might still belong to its group.
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
        await process.wait()


if __name__ == "__main__":
    try:
        # Limits are confined to this short-lived reader, never to the RAG service.
        if sys.platform == "linux":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
            resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
        raw = sys.stdin.buffer.read(MAX_FILE_BYTES + 1)
        with contextlib.redirect_stdout(sys.stderr):
            output = extract_file(raw, sys.argv[1])
    except ExtractionError as exc:
        output = {"error": str(exc), "status": exc.status}
    except ImportError:
        output = {"error": "This file reader is temporarily unavailable.", "status": 503}
    except Exception:
        logging.exception("Local file extraction failed")
        output = {"error": "The file is damaged or could not be read.", "status": 422}
    print(json.dumps(output, ensure_ascii=True))
